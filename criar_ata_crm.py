from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Users\fabri\Projetos\Escritorio de Projetos\Ata do Comitê de Marketing Modelo de Uso do CRM.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def format_table(table, widths, header_fill="244062"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9)
            else:
                if row_idx % 2 == 0:
                    set_cell_shading(cell, "F4F7FA")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.size = Pt(9)


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label)
    r.bold = True
    p.add_run(value)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor(34, 34, 34)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(24)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].paragraph_format.space_after = Pt(4)
for style_name, size in (("Heading 1", 15), ("Heading 2", 12)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(5)

title = doc.add_paragraph(style="Title")
title.add_run("Ata do Comitê de Marketing")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(10)
run = subtitle.add_run("Modelo de uso do CRM")
run.bold = True
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(62, 84, 111)

intro = doc.add_paragraph(
    "Este registro consolida as decisões do comitê sobre o modelo de uso do CRM, "
    "incluindo responsáveis, prazos, condições e pontos que exigem validação posterior."
)
intro.paragraph_format.space_after = Pt(10)

add_label_value(doc, "Data: ", "23 de setembro de 2026")
add_label_value(doc, "Horário: ", "A registrar")
add_label_value(doc, "Local ou canal: ", "A registrar")
add_label_value(doc, "Participantes: ", "A registrar")
add_label_value(doc, "Facilitador: ", "A registrar")

doc.add_heading("Objetivo da reunião", level=1)
doc.add_paragraph(
    "Definir como o CRM será utilizado pelo marketing e estabelecer regras claras para cadastro, "
    "segmentação, atualização, governança, acompanhamento e integração com outras áreas."
)

doc.add_heading("Decisões do comitê", level=1)
decisions = doc.add_table(rows=2, cols=6)
headers = ["ID", "Decisão", "Condição ou critério", "Responsável", "Prazo", "Status"]
for i, text in enumerate(headers):
    decisions.cell(0, i).text = text
values = ["D01", "", "", "", "", "Aprovada"]
for i, text in enumerate(values):
    decisions.cell(1, i).text = text
format_table(
    decisions,
    [Inches(0.42), Inches(2.25), Inches(1.72), Inches(1.12), Inches(0.82), Inches(0.8)],
)

doc.add_heading("Ações decorrentes", level=1)
actions = doc.add_table(rows=2, cols=5)
headers = ["ID", "Ação", "Responsável", "Prazo", "Situação"]
for i, text in enumerate(headers):
    actions.cell(0, i).text = text
values = ["A01", "", "", "", "Não iniciada"]
for i, text in enumerate(values):
    actions.cell(1, i).text = text
format_table(
    actions,
    [Inches(0.45), Inches(3.4), Inches(1.25), Inches(0.9), Inches(1.13)],
    header_fill="3C4A57",
)

doc.add_heading("Pontos em aberto", level=1)
open_points = doc.add_table(rows=2, cols=4)
headers = ["ID", "Questão pendente", "Quem valida", "Data limite"]
for i, text in enumerate(headers):
    open_points.cell(0, i).text = text
values = ["P01", "", "", ""]
for i, text in enumerate(values):
    open_points.cell(1, i).text = text
format_table(
    open_points,
    [Inches(0.45), Inches(4.0), Inches(1.55), Inches(1.13)],
    header_fill="5B6770",
)

doc.add_heading("Observações relevantes", level=1)
doc.add_paragraph("Registrar aqui restrições, riscos, premissas ou divergências que contextualizem as decisões.")

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("Comitê de Marketing  |  Registro de decisões sobre o CRM")
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor(100, 100, 100)

doc.core_properties.title = "Ata do Comitê de Marketing Modelo de Uso do CRM"
doc.core_properties.subject = "Registro de decisões do comitê sobre o modelo de uso do CRM"
doc.core_properties.author = "Escritório de Projetos"
doc.save(OUTPUT)
print(OUTPUT)
